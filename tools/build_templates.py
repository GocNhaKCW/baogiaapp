"""Tạo các mẫu Word có chỗ điền {…} cho app báo giá (docxtemplater).
Chạy: python3 tools/build_templates.py <thư mục chứa HDKT.docx DNTU.docx DNTT.docx> templates/
"""
import sys, re, copy, docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC, OUT = sys.argv[1], sys.argv[2]
VN = re.compile('[àáạảãâầấậẩẫăằắặẳẵèéẹẻẽêềếệểễìíịỉĩòóọỏõôồốộổỗơờớợởỡùúụủũưừứựửữỳýỵỷỹđ]', re.I)

def ptext(el): return ''.join(t.text or '' for t in el.iter(qn('w:t')))
def set_text(p, text):
    runs = p.runs
    keep = next((r for r in runs if r.text.strip()), runs[0] if runs else None)
    if keep is None: p.add_run(text); return
    for r in runs:
        if r is not keep: r._element.getparent().remove(r._element)
    keep.text = text
def remove(el): el.getparent().remove(el)
def is_en(t): return bool(t.strip()) and not VN.search(t)

# ---------------- HỢP ĐỒNG (bản tiếng Việt) ----------------
d = docx.Document(f'{SRC}/HDKT.docx')
body = d.element.body
# 1) bỏ đoạn tiếng Anh ngoài bảng
for el in list(body.iterchildren()):
    if el.tag == qn('w:p'):
        t = ptext(el)
        if is_en(t) or t.startswith('Attached to'): remove(el)
# 2) bỏ dòng tiếng Anh trong ô bảng (dòng thứ 2 trở đi)
for t in d.tables:
    for row in t.rows:
        for tc in row._tr.findall(qn('w:tc')):
            ps = tc.findall(qn('w:p'))
            for p in ps[1:]:
                if is_en(ptext(p)): remove(p)
P = {p.text: p for p in d.paragraphs}
def find(prefix):
    for p in d.paragraphs:
        if p.text.startswith(prefix): return p
    raise KeyError(prefix)
set_text(find('Số / No.'), 'Số: ……………/HĐKT/GN-……………')
set_text(find('2.1. Tổng giá trị hợp đồng'), '2.1. Tổng giá trị hợp đồng: {tong} VNĐ (đã bao gồm thuế VAT {vat_pct}% và phí vận chuyển, nếu có).')
set_text(find('(Bằng chữ:'), '(Bằng chữ: {tong_chu} đồng./.)')
p1 = find('- Lần 1:'); p2 = find('- Lần 2:')
set_text(p1, '- Lần {n}: {mota}')
st = copy.deepcopy(p1._element); en = copy.deepcopy(p1._element)
p1._element.addprevious(st); p1._element.addnext(en)
set_text(docx.text.paragraph.Paragraph(st, p1._parent), '{#dot}')
set_text(docx.text.paragraph.Paragraph(en, p1._parent), '{/dot}')
remove(p2._element)
g1 = find('- Hàng có sẵn:'); g2 = find('- Hàng đặt sản xuất')
set_text(g1, '- Đợt {n}: Giao trong vòng {ngay} ngày kể từ thời điểm Bên B nhận được khoản tạm ứng của Bên A{them}.')
st = copy.deepcopy(g1._element); en = copy.deepcopy(g1._element)
g1._element.addprevious(st); g1._element.addnext(en)
set_text(docx.text.paragraph.Paragraph(st, g1._parent), '{#giao_ds}')
set_text(docx.text.paragraph.Paragraph(en, g1._parent), '{/giao_ds}')
remove(g2._element)
set_text(find('3.2. Địa điểm giao hàng'), '3.2. Địa điểm giao hàng: Giao tại địa chỉ do Bên A chỉ định: {dcg}')
set_text(find('- Hợp đồng được lập thành'), '- Hợp đồng được lập thành 02 bản có giá trị pháp lý như nhau, mỗi bên giữ 01 bản, có hiệu lực kể từ ngày ký.')
set_text(find('Bằng chữ: …'), 'Bằng chữ: {tong_chu} đồng./.')

tA, tB, tSign, tAnx = d.tables
vals = [': {kh}', ': Ông/Bà {dd}', ': {cv}', ': {uq} ngày {uqd}', ': {dc}', ': {mst}', ': {sdt}   Email: {em}', ': {stk} tại Ngân hàng {nh}']
for row, v in zip(tA.rows, vals):
    set_text(row.cells[1].paragraphs[0], v)
set_text(tSign.rows[1].cells[0].paragraphs[0], '\n\n\n\n\n{ky_a}')
# giữ khối chữ ký trên cùng một trang
for row in tSign.rows:
    trPr = row._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); trPr.append(cs)
    for c in row.cells:
        for pp in c.paragraphs: pp.paragraph_format.keep_with_next = True
prev = tSign._tbl.getprevious()
while prev is not None and prev.tag == qn('w:p'):
    docx.text.paragraph.Paragraph(prev, d._body).paragraph_format.keep_with_next = True
    if ptext(prev).strip(): break
    prev = prev.getprevious()
# Phụ lục: 1 dòng lặp + 3 dòng tổng
rows = tAnx.rows
tpl = rows[1]
cells = tpl._tr.findall(qn('w:tc'))
keys = ['{#items}{stt}', '{ten}', '{kt}', '{mau}', '{dvt}', '{sl}', '{dg}', '{tt}', '{gc}{/items}']
for tc, k in zip(cells, keys):
    set_text(docx.table._Cell(tc, tAnx).paragraphs[0], k)
for r in rows[2:11]: remove(r._tr)
tot = tAnx.rows[-3:]
for r, (lab, val) in zip(tot, [('Cộng tiền hàng (trước thuế)', '{sau}'), ('Thuế VAT {vat_pct}%', '{vat}'), ('TỔNG CỘNG (đã bao gồm VAT)', '{tong}')]):
    tcs = r._tr.findall(qn('w:tc'))
    set_text(docx.table._Cell(tcs[0], tAnx).paragraphs[0], lab)
    set_text(docx.table._Cell(tcs[1], tAnx).paragraphs[0], val)
d.save(f'{OUT}/HDKT.docx')

# ---------------- ĐỀ NGHỊ TẠM ỨNG / THANH TOÁN (dùng chung 1 mẫu) ----------------
for name, kind in (('DNTU', 'TU'), ('DNTT', 'TT')):
    d = docx.Document(f'{SRC}/{name}.docx')
    ps = d.paragraphs
    ps[3].runs[1].text = '{kh_upper}'
    ps[4].runs[1].text = '……………………… '; ps[4].runs[3].text = '{kh} '
    if kind == 'TU':
        ps[5].runs[1].text = '{pct}% giá trị đơn hàng'; ps[5].runs[3].text = '{tien} VNĐ'
        ps[6].runs[1].text = 'Bằng chữ: {chu} đồng chẵn./.)'
    else:
        ps[5].runs[0].text = 'Chúng tôi kính đề nghị Quý Công ty thanh toán {mota} (đã bao gồm VAT), tương đương số tiền: '
        ps[5].runs[1].text = '{tien} VNĐ'
        ps[6].runs[1].text = 'Bằng chữ: {chu} đồng chẵn./.'
    set_text(ps[11], 'Nội dung: {noidung}')
    cell = d.tables[0].rows[1].cells[0].paragraphs[0]
    cell.runs[1].text = name + '-……………'
    from docx.shared import RGBColor
    for r in [ps[3].runs[1], ps[4].runs[1], ps[4].runs[3], cell.runs[1]] + list(ps[5].runs) + list(ps[6].runs) + list(ps[11].runs):
        r.font.color.rgb = RGBColor(0, 0, 0)
    d.save(f'{OUT}/{name}.docx')

# ---------------- BIÊN BẢN GIAO NHẬN (dựng lại theo mẫu) ----------------
d = docx.Document()
sec = d.sections[0]; sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(1.8)
st = d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.paragraph_format.space_after = Pt(3)
def para(text='', bold=False, align=None, size=None, italic=False, cell=None):
    p = cell.add_paragraph() if cell is not None else d.add_paragraph()
    if align: p.alignment = align
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    return p
def borders(tbl, on=True):
    tblPr = tbl._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), 'single' if on else 'nil'); x.set(qn('w:sz'), '4'); x.set(qn('w:color'), '000000'); b.append(x)
    tblPr.append(b)
C = WD_ALIGN_PARAGRAPH.CENTER
h = d.add_table(rows=1, cols=2); borders(h, False)
l, r = h.rows[0].cells
l.paragraphs[0].add_run('CÔNG TY CỔ PHẦN\nGLOBO VISTA VIỆT NAM').bold = True; l.paragraphs[0].alignment = C
para('-------------', align=C, cell=l)
r.paragraphs[0].add_run('CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM').bold = True; r.paragraphs[0].alignment = C
para('Độc lập – Tự do – Hạnh phúc', bold=True, align=C, cell=r)
para('-------------', align=C, cell=r)
para('TP Hồ Chí Minh, ngày … tháng … năm 20…', italic=True, align=C, cell=r)
para()
para('BIÊN BẢN GIAO NHẬN', bold=True, align=C, size=16)
para('Căn cứ thỏa thuận mua bán theo hợp đồng Số: ……………/HĐKT/GN-……………', italic=True, align=C)
para('Hôm nay, ngày … tháng … năm 20…, tại TP Hồ Chí Minh, chúng tôi gồm:')
para('BÊN A (Bên nhận hàng): {kh_upper}', bold=True)
para('- Địa chỉ nhận hàng: {dcg}')
para('- Đại diện: {dd}          Chức vụ: {cv}')
para('- Người nhận hàng: {nguoinhan}          Điện thoại: {sdtnhan}')
para('BÊN B (Bên giao hàng): CÔNG TY CỔ PHẦN GLOBO VISTA VIỆT NAM', bold=True)
para('Địa điểm kinh doanh số 1, Cửa hàng giới thiệu sản phẩm: Góc Nhà')
para('- Địa chỉ: 28 Vũ Tông Phan, phường Bình Trưng, TP. Hồ Chí Minh')
para('- Điện thoại: 0906 966 500')
para('- Đại diện giao hàng: ……………………………………')
para('Hai bên cùng nhau thống nhất số lượng giao hàng như sau:')
t = d.add_table(rows=2, cols=5); borders(t, True); t.alignment = WD_TABLE_ALIGNMENT.CENTER
widths = [Cm(1.2), Cm(8.8), Cm(1.6), Cm(1.8), Cm(3.6)]
for i, (hd, w) in enumerate(zip(['STT', 'Tên sản phẩm', 'ĐVT', 'Số lượng', 'Ghi chú'], widths)):
    c = t.rows[0].cells[i]; c.width = w; c.paragraphs[0].alignment = C; c.paragraphs[0].add_run(hd).bold = True
for i, (k, w) in enumerate(zip(['{#items}{stt}', '{ten}', '{dvt}', '{sl}', '{gc}{/items}'], widths)):
    c = t.rows[1].cells[i]; c.width = w; c.paragraphs[0].add_run(k)
    if i in (0, 2, 3): c.paragraphs[0].alignment = C
para()
para('Hai bên đồng ý, thống nhất ký tên. Biên bản được lập thành 02 bản, mỗi bên giữ 01 bản có giá trị pháp lý như nhau.')
para()
s = d.add_table(rows=1, cols=2); borders(s, False)
for c, txt in zip(s.rows[0].cells, ['BÊN GIAO HÀNG', 'BÊN NHẬN HÀNG']):
    c.paragraphs[0].alignment = C; c.paragraphs[0].add_run(txt).bold = True
    para('(Ký, ghi rõ họ tên)', italic=True, align=C, cell=c, size=10)
d.save(f'{OUT}/BBGH.docx')
print('ok')
